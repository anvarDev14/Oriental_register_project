import io
from datetime import datetime

from aiogram import Router, F
from aiogram.filters import Command
from aiogram.types import Message, CallbackQuery, BufferedInputFile
from aiogram.exceptions import TelegramBadRequest
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fpdf import FPDF

from app.bot import ADMIN_IDS
from app.database.requests import get_all_users, get_users_count, delete_user
from app.keyboards.admin_kb import get_admin_kb, get_back_kb

router = Router()


def is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS


# ── /admin komandasi ─────────────────────────────────────────
@router.message(Command("admin"))
async def cmd_admin(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ Sizda admin huquqi yo'q.")
        return

    count = await get_users_count()
    await message.answer(
        f"🔐 <b>Admin panel</b>\n\n"
        f"👥 Jami foydalanuvchilar: <b>{count}</b>",
        reply_markup=get_admin_kb(),
    )


# ── Admin panel callback ─────────────────────────────────────
@router.callback_query(F.data == "admin_panel")
async def admin_panel(callback: CallbackQuery):
    await callback.answer()
    if not is_admin(callback.from_user.id):
        return

    count = await get_users_count()
    try:
        await callback.message.edit_text(
            f"🔐 <b>Admin panel</b>\n\n"
            f"👥 Jami foydalanuvchilar: <b>{count}</b>",
            reply_markup=get_admin_kb(),
        )
    except TelegramBadRequest:
        pass


# ── Barcha foydalanuvchilar ro'yxati ─────────────────────────
@router.callback_query(F.data == "admin_all_users")
async def admin_all_users(callback: CallbackQuery):
    await callback.answer()
    if not is_admin(callback.from_user.id):
        return

    users = await get_all_users()

    if not users:
        try:
            await callback.message.edit_text(
                "📭 Hali hech kim ro'yxatdan o'tmagan.",
                reply_markup=get_back_kb(),
            )
        except TelegramBadRequest:
            pass
        return

    text_lines = ["👥 <b>Barcha foydalanuvchilar:</b>\n"]
    for i, user in enumerate(users, 1):
        text_lines.append(
            f"{i}. <b>{user.full_name}</b>\n"
            f"   📱 {user.phone_number}\n"
            f"   🪪 JSHSHIR: <code>{user.jshshir or '—'}</code>\n"
            f"   📄 Pasport: <b>{user.passport_id or '—'}</b>\n"
            f"   🎓 {getattr(user, 'level', 'Bakalavr')}\n"
            f"   📚 {user.direction}\n"
            f"   📋 {user.study_type}\n"
            f"   📅 {user.created_at.strftime('%d.%m.%Y %H:%M')}\n"
        )

    full_text = "\n".join(text_lines)
    if len(full_text) > 4000:
        full_text = full_text[:4000] + "\n\n<i>...ro'yxat qisqartirildi. Excel/PDF yuklab oling.</i>"

    try:
        await callback.message.edit_text(full_text, reply_markup=get_back_kb())
    except TelegramBadRequest:
        pass


# ── Statistika ───────────────────────────────────────────────
@router.callback_query(F.data == "admin_stats")
async def admin_stats(callback: CallbackQuery):
    await callback.answer()
    if not is_admin(callback.from_user.id):
        return

    users = await get_all_users()
    total = len(users)
    today = datetime.now().date()
    today_count = sum(1 for u in users if u.created_at.date() == today)

    try:
        await callback.message.edit_text(
            "📊 <b>Statistika</b>\n\n"
            f"👥 Jami ro'yxatdan o'tganlar: <b>{total}</b>\n"
            f"📅 Bugun ro'yxatdan o'tganlar: <b>{today_count}</b>",
            reply_markup=get_back_kb(),
        )
    except TelegramBadRequest:
        pass


# ── Foydalanuvchini o'chirish ────────────────────────────────
@router.callback_query(F.data.startswith("admin_delete_"))
async def admin_delete_user(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("⛔ Ruxsat yo'q.", show_alert=True)
        return

    tg_id = int(callback.data.split("_")[-1])
    deleted = await delete_user(tg_id)

    if deleted:
        await callback.answer("✅ Foydalanuvchi o'chirildi.", show_alert=True)
    else:
        await callback.answer("❌ Foydalanuvchi topilmadi.", show_alert=True)

    count = await get_users_count()
    try:
        await callback.message.edit_text(
            f"🔐 <b>Admin panel</b>\n\n"
            f"👥 Jami foydalanuvchilar: <b>{count}</b>",
            reply_markup=get_admin_kb(),
        )
    except TelegramBadRequest:
        pass


# ── Excel export ─────────────────────────────────────────────
@router.callback_query(F.data == "admin_export_excel")
async def admin_export_excel(callback: CallbackQuery):
    await callback.answer("⏳ Excel tayyorlanmoqda...")
    if not is_admin(callback.from_user.id):
        return

    users = await get_all_users()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Foydalanuvchilar"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    headers = ["#", "F.I.SH", "Telefon", "JSHSHIR", "Pasport", "Daraja", "Yo'nalish", "Shakl", "Sana"]
    col_widths = [5, 28, 16, 16, 12, 14, 40, 14, 18]

    for col, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 22

    for i, user in enumerate(users, 1):
        row_data = [
            i, user.full_name, user.phone_number,
            user.jshshir or "—", user.passport_id or "—",
            getattr(user, "level", "Bakalavr"),
            user.direction, user.study_type,
            user.created_at.strftime("%d.%m.%Y %H:%M"),
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=i + 1, column=col, value=value)
            cell.alignment = Alignment(
                horizontal="center" if col not in (2, 7) else "left", vertical="center"
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"users_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    await callback.message.answer_document(
        BufferedInputFile(buffer.read(), filename=filename),
        caption=f"📥 <b>Excel fayl</b>\n👥 Jami: <b>{len(users)}</b> ta foydalanuvchi",
    )


# ── PDF export ───────────────────────────────────────────────
@router.callback_query(F.data == "admin_export_pdf")
async def admin_export_pdf(callback: CallbackQuery):
    await callback.answer("⏳ PDF tayyorlanmoqda...")
    if not is_admin(callback.from_user.id):
        return

    users = await get_all_users()

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.add_font("DejaVu", "", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    pdf.add_font("DejaVu", "B", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")

    pdf.set_font("DejaVu", "B", 14)
    pdf.set_fill_color(46, 134, 171)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 10, "Foydalanuvchilar ro'yxati", align="C", fill=True, new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("DejaVu", "", 9)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 7, f"Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}  |  Jami: {len(users)} ta",
             align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    col_widths = [8, 46, 30, 30, 20, 18, 50, 18, 27]
    headers = ["#", "F.I.SH", "Telefon", "JSHSHIR", "Pasport", "Daraja", "Yo'nalish", "Shakl", "Sana"]

    pdf.set_font("DejaVu", "B", 9)
    pdf.set_fill_color(46, 134, 171)
    pdf.set_text_color(255, 255, 255)
    for header, width in zip(headers, col_widths):
        pdf.cell(width, 8, header, border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("DejaVu", "", 8)
    for i, user in enumerate(users):
        fill = i % 2 == 0
        pdf.set_fill_color(240, 248, 255) if fill else pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(0, 0, 0)
        row_data = [
            (str(i + 1), col_widths[0], "C"),
            (user.full_name, col_widths[1], "L"),
            (user.phone_number, col_widths[2], "C"),
            (user.jshshir or "—", col_widths[3], "C"),
            (user.passport_id or "—", col_widths[4], "C"),
            (getattr(user, "level", "Bakalavr"), col_widths[5], "C"),
            (user.direction, col_widths[6], "L"),
            (user.study_type, col_widths[7], "C"),
            (user.created_at.strftime("%d.%m.%Y %H:%M"), col_widths[8], "C"),
        ]
        for value, width, align in row_data:
            pdf.cell(width, 7, value, border=1, align=align, fill=fill)
        pdf.ln()

    buffer = io.BytesIO(pdf.output())
    filename = f"users_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    await callback.message.answer_document(
        BufferedInputFile(buffer.read(), filename=filename),
        caption=f"📄 <b>PDF fayl</b>\n👥 Jami: <b>{len(users)}</b> ta foydalanuvchi",
    )
